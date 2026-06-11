from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from build_full_report import (
    BRAND_WORDMARK,
    REPUBLICA_LOGO,
    ROOT,
    SCHOOL_LOGO,
    add_field,
    add_page_number,
    add_table,
    bullet,
    heading,
    p,
    set_styles,
    set_update_fields,
)


OUT = ROOT / "docs" / "PAP_ENTREGA" / "Manual_Tecnico_Greenerry_Srijan_Gautam.docx"


def cover(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    logo_row = doc.add_table(rows=1, cols=2)
    logo_row.alignment = WD_TABLE_ALIGNMENT.CENTER
    left = logo_row.rows[0].cells[0]
    right = logo_row.rows[0].cells[1]
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if SCHOOL_LOGO.exists():
        left.paragraphs[0].add_run().add_picture(str(SCHOOL_LOGO), width=Cm(4.7))
    if REPUBLICA_LOGO.exists():
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.paragraphs[0].add_run().add_picture(str(REPUBLICA_LOGO), width=Cm(4.2))

    for _ in range(3):
        p(doc, "", first_line=False)
    school = p(doc, "Escola Secundária Cacilhas-Tejo", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    school.runs[0].bold = True
    school.runs[0].font.size = Pt(14)
    p(doc, "Curso Profissional de Técnico de Gestão e Programação de Sistemas Informáticos", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    p(doc, "12.º K", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    for _ in range(2):
        p(doc, "", first_line=False)
    title = p(doc, "PROVA DE APTIDÃO PROFISSIONAL", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(17)
    main = p(doc, "MANUAL TÉCNICO", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    main.runs[0].bold = True
    main.runs[0].font.size = Pt(26)
    main.runs[0].font.color.rgb = RGBColor(17, 24, 39)

    if BRAND_WORDMARK.exists():
        logo = doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.add_run().add_picture(str(BRAND_WORDMARK), width=Cm(11.8))
    else:
        brand = p(doc, "GREENERRY", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        brand.runs[0].bold = True
        brand.runs[0].font.size = Pt(24)
    subtitle = p(doc, "Instalação, configuração e validação técnica", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    subtitle.runs[0].font.size = Pt(13)
    subtitle.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    for _ in range(3):
        p(doc, "", first_line=False)
    add_table(doc, ["Elemento", "Informação"], [
        ["Aluno", "Srijan Gautam"],
        ["Orientadora", "Elisabete Vaz"],
        ["Ano letivo", "2025/2026"],
    ], [4.2, 8.2])
    for _ in range(2):
        p(doc, "", first_line=False)
    date = p(doc, "Almada, junho de 2026", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    date.runs[0].font.size = Pt(11)
    date.runs[0].font.color.rgb = RGBColor(71, 85, 105)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_styles(doc)
    set_update_fields(doc)
    cover(doc)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.footer.is_linked_to_previous = False
    add_page_number(section.footer.paragraphs[0])

    heading(doc, "Índice", 1)
    toc = p(doc, first_line=False)
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Índice automático")

    heading(doc, "1. Objetivo do manual", 1)
    p(doc, "Este manual técnico explica como instalar, configurar, executar e validar a plataforma Greenerry em ambiente local. O documento foi preparado para apoiar a avaliação da Prova de Aptidão Profissional e para permitir que o projeto seja reproduzido numa máquina com XAMPP.")
    p(doc, "O manual não contém credenciais privadas de alojamento, SMTP real ou serviços externos. Apenas são apresentados dados de acesso local e credenciais de demonstração adequadas para avaliação.")

    heading(doc, "2. Requisitos de instalação", 1)
    add_table(doc, ["Requisito", "Descrição"], [
        ["Sistema operativo", "Windows, com permissões para executar XAMPP e aceder à pasta htdocs."],
        ["Servidor local", "XAMPP com Apache e MySQL/MariaDB ativos."],
        ["PHP", "Versão incluída no XAMPP usado no desenvolvimento."],
        ["Base de dados", "MySQL/MariaDB, base chamada greenerry."],
        ["Composer", "Necessário se a pasta vendor não estiver incluída."],
        ["Navegador", "Chrome, Edge, Firefox ou outro navegador moderno."],
    ], [4.2, 10.2])

    heading(doc, "3. Estrutura do projeto", 1)
    add_table(doc, ["Pasta/Ficheiro", "Função"], [
        ["admin/", "Painel administrativo: dashboard, produtos, categorias, géneros, lançamentos, encomendas, utilizadores, mensagens, relatórios, manutenção, administradores e definições."],
        ["api/", "Endpoints usados por ações assíncronas da interface."],
        ["assets/", "CSS, JavaScript, imagens, logótipos, ficheiros de áudio e recursos visuais."],
        ["includes/", "Configuração, ligação à base de dados, autenticação, validação, helpers, traduções e funções comuns."],
        ["pages/", "Páginas públicas, área de utilizador e área de artista."],
        ["vendor/", "Dependências instaladas por Composer, como PHPMailer e Dompdf."],
        ["docs/", "Relatório, manual, DER, MER/FNN e materiais de verificação PAP."],
        ["greenerry.sql", "Script de importação/criação da base de dados Greenerry."],
        ["composer.json", "Lista de dependências PHP do projeto."],
    ], [4.2, 10.2])

    heading(doc, "4. Instalação em XAMPP", 1)
    for step in [
        "Copiar a pasta greenerry para C:\\xampp\\htdocs\\dashboard\\greenerry.",
        "Abrir o painel do XAMPP.",
        "Iniciar Apache.",
        "Iniciar MySQL/MariaDB.",
        "Abrir phpMyAdmin no navegador.",
        "Criar a base de dados greenerry, caso ainda não exista.",
        "Importar o ficheiro greenerry.sql.",
        "Confirmar que a importação termina sem erros críticos.",
        "Abrir http://localhost/dashboard/greenerry/ no navegador.",
    ]:
        bullet(doc, step)

    heading(doc, "5. Configuração local", 1)
    add_table(doc, ["Campo", "Valor local"], [
        ["Host", "localhost"],
        ["Utilizador MySQL", "root"],
        ["Password MySQL", "vazia"],
        ["Base de dados", "greenerry"],
        ["Ficheiro principal de configuração", "includes/config.php"],
        ["URL local", "http://localhost/dashboard/greenerry/"],
    ], [4.5, 9.9])
    p(doc, "A configuração local foi preparada para o ambiente XAMPP. Se o projeto for colocado noutra pasta, deve ser confirmado o caminho usado no navegador e as constantes de base URL carregadas pela aplicação.")

    heading(doc, "6. Dependências", 1)
    add_table(doc, ["Dependência", "Utilização no projeto"], [
        ["PHPMailer", "Envio de e-mails de verificação, recuperação de palavra-passe, avisos e notificações."],
        ["Dompdf", "Geração de faturas/recibos em PDF."],
        ["Composer", "Gestão das dependências PHP através do ficheiro composer.json."],
    ], [4.2, 10.2])
    p(doc, "Se a pasta vendor não existir, deve ser executado o comando composer install na raiz do projeto. Depois da instalação, deve ser confirmado que o ficheiro vendor/autoload.php existe.")

    heading(doc, "7. Acessos de demonstração", 1)
    add_table(doc, ["Tipo", "Credenciais/nota"], [
        ["Admin principal", "Email: greenerry333@gmail.com | Password: Srijan123@"],
        ["Login de admin", "É feito em página reservada, apenas com e-mail e palavra-passe de uma conta administrativa ativa."],
        ["Campo extra", "Não existe campo adicional para o painel; o acesso depende da conta admin ativa."],
        ["Clientes/artistas", "O ficheiro greenerry.sql inclui dados de demonstração. As palavras-passe finais devem ser confirmadas antes da gravação do vídeo."],
    ], [4.2, 10.2])

    heading(doc, "8. Funcionalidades técnicas", 1)
    for item in [
        "Sessões PHP separadas para utilizador/artista e administração.",
        "Passwords guardadas com hash.",
        "Prepared statements em operações sensíveis.",
        "Tokens CSRF em formulários.",
        "Verificação de e-mail através de código enviado por e-mail.",
        "Recuperação de palavra-passe por código temporário.",
        "Uploads de imagens, capas e ficheiros de áudio.",
        "Carrinho, checkout, encomendas, pagamentos demonstrativos e faturas PDF.",
        "Mensagens por encomenda entre comprador e artista, funcionando como uma conversa curta ligada à compra.",
        "Notificações para utilizadores.",
        "Relatórios financeiros, musicais e de desempenho.",
        "Tema claro/escuro e idioma PT/EN.",
    ]:
        bullet(doc, item)

    heading(doc, "9. Regras de permissões", 1)
    add_table(doc, ["Perfil", "Acesso principal"], [
        ["Visitante", "Pode navegar por início, música, artistas, loja e detalhes, mas ações pessoais exigem login."],
        ["Cliente autenticado", "Pode gerir perfil, favoritos, playlists, carrinho, compras, suporte e notificações."],
        ["Artista", "Pode publicar música/produtos, consultar análises, clientes, pedidos, mensagens e rendimento."],
        ["Administrador", "Pode moderar conteúdos, gerir dados globais, responder mensagens, ver relatórios e controlar manutenção."],
    ], [4, 10.4])
    p(doc, "A moderação administrativa tem prioridade sobre as ações do artista. Quando um produto ou lançamento é bloqueado/inativado pelo administrador, o artista não consegue reativá-lo sozinho.")

    heading(doc, "10. Testes recomendados", 1)
    add_table(doc, ["Fluxo", "Teste"], [
        ["Visitante", "Abrir início, música, artistas, loja, produto e testar pesquisa."],
        ["Conta", "Registar, receber/introduzir código de e-mail e iniciar sessão."],
        ["Cliente", "Editar perfil, usar biblioteca, adicionar produto ao carrinho, finalizar checkout e consultar compras."],
        ["Comunicação", "Enviar suporte ao admin e enviar mensagem numa encomenda."],
        ["Artista", "Publicar música/produto, consultar pedidos, mensagens, análise e rendimento."],
        ["Admin", "Entrar apenas com e-mail e palavra-passe, aprovar/rejeitar/inativar conteúdos, gerir encomendas e responder mensagens."],
        ["Layout", "Verificar tema claro/escuro, mobile/desktop e logótipo em sidebar/header."],
    ], [3.6, 10.8])

    heading(doc, "11. Problemas comuns", 1)
    add_table(doc, ["Problema", "Verificação"], [
        ["Página não abre", "Confirmar Apache ativo e caminho C:\\xampp\\htdocs\\dashboard\\greenerry."],
        ["Erro de base de dados", "Confirmar que greenerry.sql foi importado e que a base se chama greenerry."],
        ["Dependências em falta", "Executar composer install e confirmar vendor/autoload.php."],
        ["E-mails não enviam", "Confirmar configuração SMTP local/de demonstração."],
        ["Imagens ou áudio não aparecem", "Confirmar caminhos dentro de assets/ e permissões de leitura."],
        ["Admin não entra", "Confirmar conta ativa na tabela admin e e-mail/password corretos."],
    ], [4.2, 10.2])

    heading(doc, "12. Materiais anexos", 1)
    p(doc, "Para a entrega PAP, este manual deve acompanhar o relatório final, o ficheiro greenerry.sql, o código do projeto, o DER, o MER/FNN, o roteiro do vídeo e a checklist de submissão.")
    add_table(doc, ["Material", "Local"], [
        ["Relatório final", "docs/PAP_ENTREGA/Relatorio_PAP_Greenerry_Srijan_Gautam_ATUALIZADO.docx"],
        ["Manual técnico", "docs/PAP_ENTREGA/Manual_Tecnico_Greenerry_Srijan_Gautam.docx"],
        ["DER", "docs/DER/DER_GREENERRY.pdf"],
        ["MER/FNN", "docs/MER_FNN/greenerry_mer_fnn.pdf"],
        ["Roteiro do vídeo", "docs/PAP_VERIFICACAO/PAP_VIDEO_ROTEIRO.md"],
        ["Checklist PAP", "docs/PAP_VERIFICACAO/PAP_REQUIREMENTS_CHECKLIST.md"],
    ], [4.4, 10])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
