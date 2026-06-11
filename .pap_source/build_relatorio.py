# -*- coding: utf-8 -*-
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"C:\xampp\htdocs\dashboard\greenerry\docs\PAP_ENTREGA")
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "Relatorio_PAP_Greenerry_Srijan_Gautam.docx"

TITLE = "Greenerry - Plataforma Web de Música, Artistas e Merchandising"
SCHOOL = "Escola Secundária Cacilhas-Tejo"
COURSE = "Curso Profissional de Técnico de Gestão e Programação de Sistemas Informáticos"
STUDENT = "Srijan Gautam"
CLASS = "12.º K"
LOCATION_DATE = "Almada, junho de 2026"

BLUE = RGBColor(20, 68, 112)
DARK = RGBColor(28, 28, 28)
MUTED = RGBColor(95, 95, 95)
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F8"
BORDER = "B7C1CC"


def set_run_font(run, name="Arial", size=11, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def paragraph_rule(paragraph, where="bottom", color="B7C1CC", size="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{where}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "4")
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)
    if placeholder:
        t = OxmlElement("w:t")
        t.text = placeholder
        run._r.append(t)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_fixed(table, widths_cm):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 567)))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_cm[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_cell_text(cell, bold=False, color=None, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.right_indent = Cm(0)
        paragraph.alignment = align
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def make_table(doc, headers, rows, widths_cm, caption=None, body_size=9.8, header_size=10.2):
    if caption:
        cap = doc.add_paragraph(caption)
        cap.style = doc.styles["Caption"]
        cap.paragraph_format.space_before = Pt(8)
        cap.paragraph_format.space_after = Pt(4)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table, widths_cm)
    set_table_borders(table)
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, LIGHT_BLUE)
        style_cell_text(cell, bold=True, color=BLUE, size=header_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header_row(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = text
            align = WD_ALIGN_PARAGRAPH.CENTER if index == len(row) - 1 and len(text) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            style_cell_text(cells[index], size=body_size, align=align)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def add_para(doc, text="", bold_start=None):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Normal"]
    if bold_start and text.startswith(bold_start):
        run = paragraph.add_run(bold_start)
        set_run_font(run, bold=True)
        rest = text[len(bold_start):]
        if rest:
            set_run_font(paragraph.add_run(rest))
    else:
        set_run_font(paragraph.add_run(text))
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.left_indent = Cm(0.9)
        paragraph.paragraph_format.hanging_indent = Cm(0.35)
        set_run_font(paragraph.add_run(item))


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0.4)
        paragraph.paragraph_format.hanging_indent = None
        set_run_font(paragraph.add_run(f"{index}. {item}"))


def add_h1(doc, text, page_break=True):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Heading 1"]
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text.upper()), size=13, bold=True, color=BLUE)


def add_h2(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Heading 2"]
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), size=12, bold=True, color=DARK)


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table, [14.0])
    set_table_borders(table, color="D4DDE6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=130, bottom=130, start=170, end=170)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(paragraph.add_run(label + ": "), size=10.5, bold=True, color=BLUE)
    set_run_font(paragraph.add_run(text), size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    for name, size, color in (("Heading 1", 13, BLUE), ("Heading 2", 12, DARK), ("Heading 3", 11.5, BLUE)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.first_line_indent = None
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(6)
    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.first_line_indent = None
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing = 1.15
    caption.paragraph_format.space_after = Pt(4)
    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(4)


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)
    try:
        section.gutter = Cm(1)
    except Exception:
        pass
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.5)
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run(TITLE), size=10, color=MUTED)
    paragraph_rule(paragraph, "bottom", color="CDD4DC", size="4")
    footer = section.footer
    f = footer.paragraphs[0]
    f.paragraph_format.first_line_indent = None
    f.paragraph_format.tab_stops.add_tab_stop(Cm(14.5), WD_ALIGN_PARAGRAPH.RIGHT)
    paragraph_rule(f, "top", color="CDD4DC", size="4")
    set_run_font(f.add_run(SCHOOL + "\tPágina "), size=10, color=MUTED)
    set_run_font(add_field(f, " PAGE ", "1"), size=10, color=MUTED)


def cover_page(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    set_run_font(p.add_run(SCHOOL), size=13, bold=True, color=DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    set_run_font(p.add_run(COURSE), size=11, color=MUTED)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    set_run_font(p.add_run("PROVA DE APTIDÃO PROFISSIONAL"), size=13, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run(TITLE), size=18, bold=True, color=DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    set_run_font(p.add_run("Relatório Final"), size=13, bold=True, color=MUTED)
    for _ in range(2):
        doc.add_paragraph()
    metadata = [
        ("Aluno", STUDENT),
        ("Turma", CLASS),
        ("Orientação", "Elisabete Vaz"),
        ("Ano letivo", "2025/2026"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table, [4.0, 8.0])
    set_table_borders(table, color="FFFFFF")
    for row, pair in zip(table.rows, metadata):
        row.cells[0].text = pair[0]
        row.cells[1].text = pair[1]
        for cell in row.cells:
            set_cell_margins(cell, top=70, bottom=70, start=100, end=100)
            style_cell_text(cell, size=11)
        style_cell_text(row.cells[0], bold=True, color=BLUE, size=11)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    set_run_font(p.add_run(LOCATION_DATE), size=11, color=MUTED)
    doc.add_page_break()


def add_toc(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(10)
    set_run_font(paragraph.add_run("ÍNDICE"), size=13, bold=True, color=BLUE)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_after = Pt(10)
    add_field(paragraph, ' TOC \\o "1-3" \\h \\z \\u ', "Atualizar índice no Word.")
    doc.add_page_break()


def add_body(doc):
    add_h1(doc, "Introdução", page_break=False)
    add_para(doc, "O presente relatório foi elaborado no âmbito da Prova de Aptidão Profissional do Curso Profissional de Técnico de Gestão e Programação de Sistemas Informáticos, da turma 12.º K, na Escola Secundária Cacilhas-Tejo. O projeto desenvolvido tem o nome Greenerry e consiste numa plataforma Web de música, artistas e merchandising, criada com o objetivo de juntar, num único sistema, a descoberta musical, a gestão de artistas, a loja online e uma área administrativa completa.")
    add_para(doc, "A ideia principal do projeto surgiu da necessidade de criar uma aplicação prática, com ligação real a base de dados, autenticação, gestão de conteúdos, funcionalidades de utilizador e processos de administração. A Greenerry permite que um visitante descubra músicas e artistas, que um utilizador autenticado guarde favoritos e realize compras, que artistas façam a gestão dos seus conteúdos e que a administração controle produtos, géneros, lançamentos, encomendas, mensagens e configurações gerais da plataforma.")
    add_para(doc, "Ao longo do desenvolvimento, o projeto foi tratado como uma aplicação completa e não apenas como um conjunto de páginas isoladas. Por esse motivo, foram trabalhadas áreas como estrutura da base de dados, consultas SQL, organização de código, segurança básica, validação de formulários, interface responsiva, envio de e-mails, geração de documentos PDF e construção de um leitor de música funcional.")
    add_callout(doc, "Tema da PAP", "Greenerry - Plataforma Web de Música, Artistas e Merchandising.")

    add_h1(doc, "1. Identificação do problema / necessidade")
    add_para(doc, "Atualmente, muitos artistas independentes precisam de várias ferramentas diferentes para divulgar músicas, mostrar lançamentos, vender produtos e comunicar com o público. Esta separação torna a gestão mais difícil e também torna a experiência do utilizador menos direta, porque o fã precisa de procurar informações em vários locais diferentes.")
    add_para(doc, "A Greenerry responde a essa necessidade através de uma plataforma única, onde a parte musical e a parte comercial funcionam em conjunto. O utilizador pode descobrir artistas, ouvir músicas, seguir perfis, consultar produtos de merchandising, adicionar itens ao carrinho e acompanhar encomendas. Por outro lado, a administração consegue gerir conteúdos, moderar publicações, acompanhar mensagens e manter a plataforma organizada.")
    add_para(doc, "O problema técnico também foi relevante: era necessário construir um sistema suficientemente completo para demonstrar conhecimentos de programação, base de dados, autenticação, organização de interfaces e integração entre várias partes do website. Assim, a PAP não ficou limitada a uma página visual, mas sim a uma aplicação com diferentes papéis de utilizador e várias operações reais.")
    add_h2(doc, "Público-alvo")
    add_bullets(doc, [
        "Utilizadores que procuram descobrir músicas, artistas e produtos relacionados com música.",
        "Artistas que pretendem divulgar lançamentos e acompanhar parte da sua atividade na plataforma.",
        "Administradores responsáveis por gerir conteúdos, encomendas, mensagens, géneros, produtos e utilizadores.",
    ])

    add_h1(doc, "2. Objetivos do projeto")
    add_h2(doc, "Objetivo geral")
    add_para(doc, "O objetivo geral da Greenerry foi desenvolver uma plataforma Web funcional, responsiva e ligada a uma base de dados, capaz de integrar música, perfis de artistas, loja online e painel administrativo numa experiência coerente.")
    add_h2(doc, "Objetivos específicos")
    add_bullets(doc, [
        "Criar um sistema de registo, autenticação, verificação de e-mail e recuperação de palavra-passe.",
        "Desenvolver páginas públicas para música, artistas, produtos e detalhes de conteúdos.",
        "Implementar um leitor de música com fila de reprodução, informações da faixa e controlos básicos.",
        "Criar favoritos, biblioteca pessoal, playlists e seguimento de artistas.",
        "Construir uma loja online com produtos, tamanhos, stock, carrinho, checkout, encomendas e recibos.",
        "Criar uma área de artista para gerir uploads, produtos, lançamentos, encomendas e estatísticas.",
        "Criar um painel administrativo para gerir utilizadores, música, géneros, lançamentos, produtos, categorias, encomendas, mensagens e configurações.",
        "Garantir organização de base de dados com relações entre tabelas e consultas SQL adequadas.",
        "Aplicar cuidados de segurança, como sessões, validação, permissões e consultas preparadas.",
        "Produzir documentação final, incluindo relatório, manual técnico, manual de utilização e guião de vídeo.",
    ])

    add_h1(doc, "3. Metodologias de trabalho")
    add_para(doc, "A metodologia utilizada foi incremental. O desenvolvimento começou pela análise das necessidades principais da aplicação, seguindo-se o planeamento da estrutura das páginas, da base de dados e dos diferentes tipos de utilizador. Depois, o projeto foi dividido por áreas funcionais, o que facilitou a implementação gradual e os testes em ambiente local.")
    add_numbered(doc, [
        "Levantamento de requisitos e definição das funcionalidades principais da plataforma.",
        "Planeamento da estrutura da base de dados, incluindo entidades, relações e campos essenciais.",
        "Criação das páginas principais e definição de uma identidade visual coerente.",
        "Implementação da autenticação, permissões e áreas privadas.",
        "Desenvolvimento da loja, do sistema de música, do painel administrativo e da área de artista.",
        "Testes locais com XAMPP, correção de erros, ajustes de layout e melhoria da experiência de utilização.",
        "Organização dos anexos, diagramas, documentação técnica e material para apresentação.",
    ])
    add_para(doc, "Durante o processo, sempre que era identificada uma falha, a solução era testada primeiro em ambiente local. Esta forma de trabalho permitiu corrigir erros de base de dados, problemas de layout, comportamentos do leitor de música e falhas em formulários sem afetar o restante projeto.")

    add_h1(doc, "4. Ferramentas e tecnologias utilizadas")
    add_para(doc, "A escolha das tecnologias teve em conta o contexto do curso e a necessidade de criar uma aplicação Web completa, executada localmente com XAMPP e suportada por uma base de dados MySQL/MariaDB.")
    make_table(doc, ["Tecnologia", "Tipo", "Utilização no projeto"], [
        ("HTML", "Frontend", "Estrutura das páginas, formulários, listas, botões, cartões e componentes visuais."),
        ("CSS", "Frontend", "Definição do layout, cores, responsividade, painéis, barras laterais, tabelas e adaptação a diferentes ecrãs."),
        ("JavaScript", "Frontend", "Interações da interface, leitor de música, menus laterais, pesquisa dinâmica e comportamentos visuais."),
        ("PHP", "Backend", "Processamento de pedidos, sessões, permissões, formulários, consultas à base de dados e geração de respostas."),
        ("MySQL/MariaDB", "Base de dados", "Armazenamento de utilizadores, faixas, géneros, lançamentos, produtos, encomendas, mensagens e notificações."),
        ("XAMPP", "Ambiente local", "Execução local do Apache, PHP e MySQL durante o desenvolvimento e testes."),
        ("phpMyAdmin", "Gestão da BD", "Consulta, importação, verificação e manutenção da base de dados."),
        ("Composer", "Dependências", "Gestão das bibliotecas externas utilizadas no projeto."),
        ("PHPMailer", "E-mail", "Envio de e-mails de verificação, recuperação de palavra-passe e comunicação do sistema."),
        ("Dompdf", "PDF", "Geração de documentos PDF, como recibos ou faturas associados a encomendas."),
        ("diagrams.net", "Modelação", "Construção dos diagramas MER, DER e FNN usados na documentação."),
    ], [3.8, 3.0, 7.2], "Tabela 1 - Ferramentas e tecnologias utilizadas.")
    add_para(doc, "Estas ferramentas permitiram desenvolver uma solução coerente com os conteúdos aprendidos ao longo do curso, principalmente nas áreas de programação Web, bases de dados, modelação de dados e desenvolvimento de aplicações com várias camadas.")

    add_h1(doc, "5. Estrutura da base de dados")
    add_para(doc, "A base de dados é uma das partes mais importantes do projeto, porque suporta praticamente todas as funcionalidades da Greenerry. Foi necessário organizar entidades ligadas a utilizadores, administradores, música, géneros, produtos, stock, encomendas, pagamentos, mensagens, notificações e segurança de conta.")
    add_para(doc, "A estrutura foi pensada para reduzir repetição de dados e permitir relações claras entre as tabelas. Por exemplo, uma faixa pertence a um género e pode estar associada a um lançamento musical; um produto pertence a uma categoria e pode ter várias imagens; uma encomenda pertence a um cliente e contém vários itens.")
    add_h2(doc, "Tabelas principais")
    make_table(doc, ["Tabela", "Finalidade"], [
        ("admin", "Guarda as contas administrativas e os níveis de acesso usados no painel de administração."),
        ("cliente", "Armazena utilizadores, artistas e dados de autenticação, perfil e estado da conta."),
        ("categoria", "Organiza os produtos da loja por categorias comerciais."),
        ("genero", "Guarda os géneros musicais usados para classificar faixas e lançamentos."),
        ("release_musical", "Regista lançamentos musicais, como álbuns, singles ou projetos associados a artistas."),
        ("faixa", "Contém as músicas, ficheiros de áudio, capas, contadores e relações com artistas e géneros."),
        ("produto", "Armazena produtos de merchandising, preços, descrição, estado e relação com vendedores."),
        ("produto_imagem", "Permite associar várias imagens a cada produto."),
        ("produto_tamanho_stock", "Controla tamanhos e quantidades disponíveis em stock."),
        ("encomenda", "Regista compras feitas pelos clientes, estados, totais e informação geral do pedido."),
        ("encomenda_item", "Guarda os produtos individuais incluídos em cada encomenda."),
        ("pagamento", "Regista informação de pagamento associada às encomendas."),
        ("favorito_musica", "Permite que o utilizador guarde músicas na sua biblioteca ou favoritos."),
        ("playlist", "Guarda listas de reprodução criadas pelos utilizadores."),
        ("playlist_faixa", "Relaciona playlists com as faixas incluídas em cada uma."),
        ("faixa_listen", "Regista reproduções de faixas para estatísticas e contagem de atividade."),
        ("seguir_artista", "Permite que utilizadores sigam artistas dentro da plataforma."),
        ("mensagem_admin", "Guarda mensagens enviadas entre utilizadores/artistas e a administração."),
        ("notificacao", "Regista avisos internos enviados para utilizadores."),
        ("recuperacao_password", "Guarda pedidos temporários de recuperação de palavra-passe."),
        ("verificacao_email", "Guarda códigos e tokens de verificação de e-mail."),
    ], [4.6, 9.4], "Tabela 2 - Principais tabelas da base de dados Greenerry.", body_size=9.5)
    add_h2(doc, "Relações importantes")
    add_bullets(doc, [
        "Um cliente pode ter várias encomendas, várias playlists, várias músicas favoritas e vários artistas seguidos.",
        "Um produto pertence a uma categoria e pode ter imagens e registos de stock por tamanho.",
        "Uma encomenda contém vários itens, permitindo guardar diferentes produtos dentro da mesma compra.",
        "Uma faixa está associada a um artista, a um género e, quando aplicável, a um lançamento musical.",
        "Um género pode classificar várias faixas e lançamentos, permitindo pesquisa, filtros e estatísticas por estilo musical.",
        "As tabelas de recuperação de palavra-passe e verificação de e-mail ajudam a separar dados temporários da tabela principal de clientes.",
    ])
    add_callout(doc, "Nota", "O diagrama MER/FNN e o DER fazem parte dos anexos do projeto, servindo para representar visualmente as entidades, relações e normalização da base de dados.")

    add_h1(doc, "6. Desenvolvimento do projeto")
    add_h2(doc, "Estrutura geral da aplicação")
    add_para(doc, "O projeto foi organizado por pastas, separando as páginas públicas, a área administrativa, os pedidos de API, os ficheiros de apoio e os recursos visuais. Esta divisão ajudou a manter o código mais organizado e facilitou a manutenção das funcionalidades.")
    add_bullets(doc, [
        "A pasta admin contém o painel administrativo e as páginas de gestão interna.",
        "A pasta pages contém as páginas principais acessíveis aos utilizadores da plataforma.",
        "A pasta api contém pontos de comunicação usados por funcionalidades dinâmicas, como música e interações assíncronas.",
        "A pasta includes reúne ficheiros comuns, como ligação à base de dados, funções auxiliares, autenticação, e-mail e regras de permissões.",
        "A pasta assets contém CSS, JavaScript, imagens, capas, ficheiros de áudio e outros recursos usados pela interface.",
        "A pasta docs contém documentação, diagramas e ficheiros de apoio para a PAP.",
    ])
    add_h2(doc, "Frontend")
    add_para(doc, "No frontend, o objetivo foi criar uma interface moderna, escura, coerente e funcional. Foram trabalhadas páginas de música, artistas, loja, detalhes de produto, carrinho, checkout, área de utilizador, área de artista e painel administrativo. Também foram criados componentes comuns, como barras laterais, cartões de música, tabelas, formulários, botões, pesquisa e filtros.")
    add_para(doc, "A responsividade foi uma preocupação importante. A interface teve de funcionar em ecrãs grandes e pequenos, sem sobreposição de elementos, com painéis laterais adaptados e conteúdos organizados em grelhas que se ajustam ao espaço disponível.")
    add_h2(doc, "Backend")
    add_para(doc, "No backend, o PHP foi usado para receber formulários, validar dados, gerir sessões, consultar a base de dados, aplicar permissões e devolver respostas às páginas. Foram usadas consultas preparadas para reduzir riscos de injeção SQL e para garantir maior segurança no acesso aos dados.")
    add_para(doc, "A autenticação foi separada por perfis, permitindo distinguir utilizadores comuns, artistas e administradores. Esta separação foi essencial para que cada área mostrasse apenas as ações correspondentes ao papel de cada utilizador.")
    add_h2(doc, "Leitor de música")
    add_para(doc, "O leitor de música foi uma das partes mais exigentes. Foi necessário controlar a faixa atual, a fila de reprodução, a capa, o artista, os botões de reprodução e a persistência do estado visual entre páginas. O objetivo foi aproximar a experiência daquilo que se espera de uma plataforma musical, mantendo o leitor disponível no rodapé da aplicação.")
    add_h2(doc, "Painel administrativo")
    add_para(doc, "O painel administrativo foi desenvolvido para permitir a gestão interna da plataforma. A administração pode consultar indicadores, gerir música, lançamentos, géneros, produtos, categorias, encomendas, utilizadores, mensagens, relatórios e definições. Esta área é essencial para manter os conteúdos controlados e garantir que a plataforma pode ser administrada sem mexer diretamente na base de dados.")

    add_h1(doc, "7. Funcionalidades implementadas e extras")
    add_para(doc, "A Greenerry inclui funcionalidades divididas por três grandes áreas: utilizador, artista e administração. Esta divisão tornou o projeto mais completo e permitiu demonstrar diferentes tipos de operações sobre a base de dados.")
    add_h2(doc, "Funcionalidades do utilizador")
    add_bullets(doc, [
        "Registo, login, logout, verificação de e-mail e recuperação de palavra-passe.",
        "Consulta de músicas, artistas, lançamentos e produtos.",
        "Pesquisa e navegação por conteúdos musicais e comerciais.",
        "Leitor de música com apresentação de faixa, artista, capa e fila.",
        "Favoritos, biblioteca pessoal, playlists e seguimento de artistas.",
        "Carrinho de compras, checkout, histórico de encomendas e recibos/faturas em PDF.",
        "Perfil de utilizador, moradas, notificações e mensagens de apoio.",
    ])
    add_h2(doc, "Funcionalidades do artista")
    add_bullets(doc, [
        "Gestão de perfil artístico e presença na plataforma.",
        "Envio e organização de músicas e lançamentos.",
        "Gestão de produtos de merchandising associados ao artista.",
        "Consulta de encomendas relacionadas com produtos do artista.",
        "Acompanhamento de estatísticas e indicadores de atividade.",
    ])
    add_h2(doc, "Funcionalidades da administração")
    add_bullets(doc, [
        "Dashboard com indicadores gerais da plataforma.",
        "Gestão de utilizadores, artistas e permissões.",
        "Gestão de produtos, categorias, tamanhos, stock e imagens.",
        "Gestão de géneros musicais, incluindo criação, edição, listagem e estado ativo/inativo.",
        "Gestão de músicas, lançamentos e moderação de conteúdos.",
        "Gestão de encomendas, mensagens, relatórios, notificações e definições do site.",
        "Suporte a idioma português/inglês e alternância de tema claro/escuro.",
    ])
    add_h2(doc, "Extras valorizáveis")
    add_para(doc, "Como extras, foram adicionados envio de e-mails, geração de PDFs, área de artista, relatórios, painel administrativo completo, notificações, gestão de géneros e melhorias de layout.")

    add_h1(doc, "8. Resultados obtidos")
    add_para(doc, "O resultado final é uma plataforma Web funcional, com várias áreas interligadas e suportadas por uma base de dados estruturada. A Greenerry permite navegar por músicas, artistas e produtos, criar conta, confirmar e-mail, utilizar funcionalidades pessoais, efetuar compras simuladas e administrar conteúdos através de um painel próprio.")
    add_para(doc, "Do ponto de vista técnico, o projeto demonstra a utilização de PHP, MySQL, HTML, CSS e JavaScript num sistema com autenticação, relações entre tabelas, formulários, uploads, permissões, consultas SQL e geração de documentos. A existência de uma área administrativa e de uma área de artista aumenta a complexidade do projeto e aproxima-o de uma aplicação real.")
    add_para(doc, "A parte visual também foi trabalhada para apresentar uma identidade própria. O layout escuro, os cartões de música, o leitor no rodapé, as barras laterais e o painel administrativo ajudam a criar uma experiência mais consistente e adequada ao tema musical.")
    make_table(doc, ["Área testada", "Resultado esperado", "Estado"], [
        ("Autenticação", "Registo, login, logout, verificação de e-mail e recuperação de palavra-passe.", "Funcional"),
        ("Música", "Listagem de faixas, géneros, lançamentos, favoritos e leitor.", "Funcional"),
        ("Loja", "Produtos, tamanhos, stock, carrinho, checkout e encomendas.", "Funcional"),
        ("Administração", "Gestão de conteúdos, utilizadores, produtos, géneros, encomendas e mensagens.", "Funcional"),
        ("Área de artista", "Gestão de uploads, produtos, encomendas e estatísticas.", "Funcional"),
        ("Documentos", "Geração de recibos/faturas e preparação de anexos da PAP.", "Funcional"),
    ], [3.4, 8.0, 2.6], "Tabela 3 - Síntese dos testes funcionais realizados.")

    add_h1(doc, "9. Dificuldades sentidas")
    add_para(doc, "Durante o desenvolvimento da Greenerry surgiram várias dificuldades, sobretudo por se tratar de uma plataforma com muitas áreas diferentes. A complexidade aumentou porque a música, a loja, o painel administrativo, a área de artista e o sistema de utilizadores tinham de funcionar em conjunto.")
    add_h2(doc, "Base de dados")
    add_para(doc, "A base de dados foi uma das maiores dificuldades. Foi necessário pensar nas tabelas, nas relações entre entidades e nas consultas SQL necessárias para apresentar informação correta em cada página. Como existiam produtos, faixas, géneros, lançamentos, encomendas, utilizadores, favoritos, playlists e notificações, qualquer alteração numa tabela podia afetar várias páginas do projeto.")
    add_h2(doc, "Consultas SQL")
    add_para(doc, "Algumas consultas exigiram junções entre várias tabelas e filtros por estado, utilizador, artista, género ou encomenda. Também foi necessário garantir que os resultados apareciam de forma correta no painel administrativo e nas páginas públicas, sem duplicar dados nem perder informação importante.")
    add_h2(doc, "Leitor de música")
    add_para(doc, "Fazer o leitor de música funcionar corretamente também foi difícil. O leitor precisava de mostrar a música atual, manter a capa e o artista corretos, lidar com a fila de reprodução e continuar integrado com a interface geral. Pequenos erros de JavaScript podiam causar falhas nos controlos ou no estado visual do leitor.")
    add_h2(doc, "Gestão do layout")
    add_para(doc, "A gestão do layout exigiu muitos ajustes. Algumas páginas tinham cartões, tabelas, barras laterais, formulários e painéis que precisavam de se adaptar a diferentes resoluções. Foi necessário corrigir espaçamentos, tamanhos de componentes, comportamento das barras laterais e encaixe de conteúdos, principalmente no painel administrativo e no leitor.")
    add_h2(doc, "Complexidade geral")
    add_para(doc, "Outra dificuldade foi manter o projeto organizado enquanto novas funcionalidades eram adicionadas. Como a plataforma tem muitas páginas e diferentes permissões, foi importante evitar que uma alteração numa parte prejudicasse outra parte do website.")

    add_h1(doc, "10. Resolução das dificuldades")
    add_para(doc, "As dificuldades foram resolvidas através de pesquisa, testes sucessivos e reorganização gradual do projeto. Em vez de tentar resolver todos os problemas ao mesmo tempo, cada área foi analisada separadamente até funcionar de forma estável.")
    add_h2(doc, "Base de dados e consultas")
    add_para(doc, "Para resolver os problemas da base de dados, foram criadas tabelas com responsabilidades bem definidas e relações adequadas. A utilização de chaves estrangeiras e a separação de dados por entidades ajudaram a reduzir repetição e facilitar consultas. As consultas SQL foram ajustadas progressivamente, testando primeiro os resultados e só depois integrando nas páginas.")
    add_h2(doc, "Leitor de música")
    add_para(doc, "No leitor de música, a solução passou por separar melhor o estado da interface, os dados da faixa e os controlos de reprodução. A fila de músicas e os dados apresentados ao utilizador foram testados em diferentes páginas para garantir que o leitor continuava coerente.")
    add_h2(doc, "Layout e interface")
    add_para(doc, "Na parte visual, foram feitos ajustes de CSS e organização dos componentes. As barras laterais foram revistas, os cartões foram ajustados para encaixar melhor os conteúdos e o painel administrativo foi melhorado para apresentar os géneros e outros dados de forma mais limpa e legível.")
    add_h2(doc, "Organização do código")
    add_para(doc, "A organização por pastas e ficheiros comuns ajudou a evitar repetição. Ficheiros partilhados, como ligação à base de dados, autenticação, funções auxiliares e cabeçalhos, permitiram manter uma estrutura mais estável e facilitar futuras alterações.")

    add_h1(doc, "11. Conclusão")
    add_para(doc, "A realização da PAP permitiu aplicar conhecimentos adquiridos ao longo do curso de Gestão e Programação de Sistemas Informáticos num projeto prático e completo. A Greenerry não se limitou a demonstrar páginas visuais, mas sim uma aplicação com base de dados, autenticação, permissões, painel administrativo, loja, música, área de artista e documentação.")
    add_para(doc, "Durante o desenvolvimento, foram consolidados conhecimentos de PHP, MySQL, HTML, CSS, JavaScript, organização de ficheiros, consultas SQL, validação de formulários e integração de bibliotecas externas. As maiores aprendizagens surgiram precisamente nas partes mais difíceis, como a estrutura da base de dados, o leitor de música, o layout e a ligação entre várias áreas do website.")
    add_para(doc, "O projeto final representa uma evolução importante nas capacidades técnicas e na autonomia de desenvolvimento. Também mostrou a importância de planear antes de programar, testar com frequência, documentar o trabalho e manter uma visão global do sistema.")

    add_h1(doc, "12. Melhorias futuras")
    add_para(doc, "Apesar de a plataforma cumprir os objetivos principais definidos para a PAP, existem várias melhorias que poderiam ser desenvolvidas futuramente para aproximar ainda mais a Greenerry de uma aplicação pronta para produção.")
    add_bullets(doc, [
        "Integração com pagamentos reais, como MB WAY, cartão ou PayPal.",
        "Sistema de recomendações musicais baseado em géneros, favoritos e histórico de audição.",
        "Chat em tempo real entre utilizadores, artistas e administração.",
        "Aplicação mobile ou versão progressiva para telemóveis.",
        "Melhoria das estatísticas para artistas, incluindo gráficos mais completos.",
        "Maior automatização da moderação de conteúdos enviados por artistas.",
        "Melhorias de acessibilidade, segurança, testes automáticos e desempenho.",
        "Integração com serviços externos de streaming ou distribuição musical.",
    ])

    add_h1(doc, "13. Referências bibliográficas / webgrafia")
    add_para(doc, "As referências seguintes foram utilizadas como apoio técnico durante o desenvolvimento e documentação do projeto:")
    add_bullets(doc, [
        "PHP Manual. Documentação oficial da linguagem PHP. Disponível em: https://www.php.net/manual/",
        "MDN Web Docs. Documentação de HTML, CSS e JavaScript. Disponível em: https://developer.mozilla.org/",
        "MySQL Documentation. Documentação oficial do MySQL. Disponível em: https://dev.mysql.com/doc/",
        "MariaDB Documentation. Documentação oficial do MariaDB. Disponível em: https://mariadb.com/kb/en/documentation/",
        "Apache Friends. Documentação e informação sobre XAMPP. Disponível em: https://www.apachefriends.org/",
        "PHPMailer. Biblioteca para envio de e-mails em PHP. Disponível em: https://github.com/PHPMailer/PHPMailer",
        "Dompdf. Biblioteca para geração de documentos PDF em PHP. Disponível em: https://github.com/dompdf/dompdf",
        "diagrams.net. Ferramenta usada para criação de diagramas. Disponível em: https://www.diagrams.net/",
    ])

    add_h1(doc, "14. Anexos")
    add_para(doc, "Os anexos complementam o relatório e devem acompanhar a entrega final da PAP. Estes materiais ajudam a demonstrar a estrutura técnica, o funcionamento visual e a preparação da apresentação.")
    make_table(doc, ["Anexo", "Conteúdo"], [
        ("Anexo A", "Diagrama MER/FNN da base de dados Greenerry."),
        ("Anexo B", "DER da base de dados e estrutura relacional."),
        ("Anexo C", "Capturas de ecrã da página inicial, música, loja, carrinho, área de artista e painel administrativo."),
        ("Anexo D", "Manual técnico com instalação, configuração, base de dados e estrutura de ficheiros."),
        ("Anexo E", "Manual de utilização para utilizador, artista e administrador."),
        ("Anexo F", "Guião do vídeo/manual de apresentação."),
    ], [3.0, 11.0], "Tabela 4 - Lista de anexos previstos para a entrega.")
    add_para(doc, "Na entrega final, os anexos devem ser impressos ou incluídos de acordo com as orientações da escola, juntamente com o relatório, o manual técnico, o manual de utilização, o vídeo/manual e o software funcional.")


def audit_text(doc_path):
    with zipfile.ZipFile(doc_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
    bad_markers = ["Ã£", "Ã¡", "Ã©", "Ãª", "Ã­", "Ã³", "Ãº", "Ã§", "Âº", "Âª", "�", "â€™", "â€œ", "â€"]
    return [marker for marker in bad_markers if marker in document_xml]


def main():
    doc = Document()
    setup_styles(doc)
    setup_page(doc)
    cover_page(doc)
    add_toc(doc)
    add_body(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.author = STUDENT
    doc.core_properties.subject = "Relatório da Prova de Aptidão Profissional"
    doc.core_properties.keywords = "PAP, Greenerry, GPSI, música, PHP, MySQL"
    doc.save(DOCX)
    bad_markers = audit_text(DOCX)
    if bad_markers:
        print("WARN bad markers:", bad_markers)
    else:
        print("OK", DOCX)


if __name__ == "__main__":
    main()
